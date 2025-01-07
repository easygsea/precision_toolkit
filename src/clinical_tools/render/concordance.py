from tabulate import tabulate
from docx import Document
from clinical_tools.analysis.compute_variant_metrics import process_variant_metrics


def save_concordance_markdown(results, output_file):
  """Save variant concordance results in Markdown format."""
  headers = [
    "Variant Type", "Variant Filter", "Concordance", "CI Lower", "CI Upper"
  ]
  data = [
    [variant_type, variant_filter, f"{concordance:.4f}", f"{ci_lower:.4f}",
     f"{ci_upper:.4f}"]
    for (variant_type, variant_filter),
        (concordance, ci_lower, ci_upper) in results.items()
  ]

  markdown_table = tabulate(data, headers, tablefmt="github")

  summary = (
    "The table below summarizes the concordance metrics calculated for each "
    "variant type and filter."
  )

  with open(output_file, "w") as f:
    f.write("# Variant Concordance\n\n")
    f.write(f"{summary}\n\n")
    f.write(markdown_table)
  print(f"Variant concordance results saved to {output_file}")


def save_concordance_word(results, output_file):
  """Save variant concordance results in Word format."""
  doc = Document()
  doc.add_heading("Variant Concordance", level=1)

  summary = (
    "The table below summarizes the concordance metrics calculated for each "
    "variant type and filter."
  )
  doc.add_paragraph(summary)

  headers = [
    "Variant Type", "Variant Filter", "Concordance", "CI Lower", "CI Upper"
  ]
  table = doc.add_table(rows=len(results) + 1, cols=len(headers))

  for j, header in enumerate(headers):
    table.cell(0, j).text = header


  for i, ((variant_type, variant_filter),
          (concordance, ci_lower, ci_upper)) in enumerate(
              results.items(), start=1):
    table.cell(i, 0).text = variant_type
    table.cell(i, 1).text = variant_filter
    table.cell(i, 2).text = f"{concordance:.4f}"
    table.cell(i, 3).text = f"{ci_lower:.4f}"
    table.cell(i, 4).text = f"{ci_upper:.4f}"

  doc.save(output_file)
  print(f"Variant concordance results saved to {output_file}")
